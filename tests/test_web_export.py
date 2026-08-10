"""The export downloads (specification 8)."""

from io import BytesIO

from openpyxl import load_workbook


def test_markdown_wird_als_download_geliefert(client, graph):
    antwort = client.get("/export/markdown")
    assert antwort.status_code == 200
    assert antwort.headers["content-type"].startswith("text/markdown")
    assert "attachment" in antwort.headers["content-disposition"]
    assert "### BFS 26a — Deutsch" in antwort.text


def test_xlsx_wird_als_download_geliefert(client, graph):
    antwort = client.get("/export/xlsx")
    assert antwort.status_code == 200
    assert "spreadsheetml" in antwort.headers["content-type"]
    assert "attachment" in antwort.headers["content-disposition"]

    mappe = load_workbook(BytesIO(antwort.content))
    assert mappe.sheetnames == ["BFS 26a Deutsch"]


def test_der_dateiname_enthaelt_keine_gespeicherten_daten(client, session, graph):
    """Nothing from an Excel import can break the header."""
    graph.klasse.bezeichnung = 'Klasse "mit; Sonderzeichen"'
    session.commit()

    kopfzeile = client.get("/export/xlsx").headers["content-disposition"]
    assert "Sonderzeichen" not in kopfzeile
    assert kopfzeile.startswith('attachment; filename="notenverwaltung-export-')
    assert kopfzeile.isascii()


def test_export_ohne_daten_liefert_trotzdem_eine_datei(client):
    assert client.get("/export/markdown").status_code == 200
    assert client.get("/export/xlsx").status_code == 200


def test_die_verwaltung_verlinkt_beide_exporte(client):
    antwort = client.get("/verwaltung")
    assert '/export/xlsx' in antwort.text
    assert '/export/markdown' in antwort.text
    assert "Klarnamen" in antwort.text
