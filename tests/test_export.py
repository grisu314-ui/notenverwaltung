"""Export of the whole dataset (specification 8)."""

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.db.models import Klasse, Kurs, Leistung, Note, Notengruppe
from app.enums import NoteStatus
from app.services.export import (
    NICHT_ERBRACHT,
    NICHT_GEWERTET,
    OHNE_NOTE,
    ZEILE_ERSTER_SCHUELER,
    ZEILE_LEISTUNG,
    als_markdown,
    als_xlsx,
    blattname,
)
from app.services.kursblatt import ANSICHT_JAHR, blatt


def geladen(rohdaten: bytes):
    return load_workbook(BytesIO(rohdaten))


def test_markdown_hat_einen_abschnitt_je_kurs(session, graph):
    text = als_markdown(session)
    assert "## Schuljahr 2026/27" in text
    assert "### BFS 26a — Deutsch" in text
    assert "Öztürk, Änne" in text
    assert "1. Klassenarbeit (15.09.2026)" in text


def test_markdown_zeigt_noten_als_zeichen(session, graph):
    graph.note.notenwert = Decimal("0.7")
    session.commit()

    text = als_markdown(session)
    assert "| 1+ |" in text
    assert "0.7" not in text


def test_markdown_kennzeichnet_die_sonderstatus(session, graph):
    session.add(
        Note(
            leistung=graph.leistung,
            schueler=graph.schueler_b,
            notenwert=None,
            status=NoteStatus.NICHT_ERBRACHT,
        )
    )
    session.commit()

    text = als_markdown(session)
    assert NICHT_ERBRACHT in text
    assert "zählt als 6" in text  # Legende


def test_markdown_ohne_kurse_ist_trotzdem_gueltig(session):
    text = als_markdown(session)
    assert "Es ist noch kein Kurs angelegt." in text


def test_markdown_weist_auf_klarnamen_hin(session, graph):
    """Specification 8 and 11: the operator has to know what is in the file."""
    assert "Klarnamen" in als_markdown(session)


def test_xlsx_hat_ein_blatt_je_kurs(session, graph):
    session.add(Kurs(klasse=graph.klasse, fach="Wirtschaftslehre"))
    session.commit()

    mappe = geladen(als_xlsx(session))
    assert sorted(mappe.sheetnames) == ["BFS 26a Deutsch", "BFS 26a Wirtschaftslehre"]


def test_xlsx_kopfbereich_nennt_klasse_kurs_und_schuljahr(session, graph):
    blattobjekt = geladen(als_xlsx(session))["BFS 26a Deutsch"]
    assert blattobjekt.cell(row=1, column=1).value == "BFS 26a — Deutsch"
    assert "Schuljahr 2026/27" in blattobjekt.cell(row=2, column=1).value


def test_xlsx_beschriftungen_sind_gedreht(session, graph):
    """Open point O-6: the label cells are rotated 90 degrees."""
    blattobjekt = geladen(als_xlsx(session))["BFS 26a Deutsch"]
    assert blattobjekt.cell(row=ZEILE_LEISTUNG, column=2).alignment.textRotation == 90


def test_xlsx_hat_eine_zeile_je_schueler(session, graph):
    blattobjekt = geladen(als_xlsx(session))["BFS 26a Deutsch"]
    namen = [
        blattobjekt.cell(row=zeile, column=1).value
        for zeile in range(ZEILE_ERSTER_SCHUELER, ZEILE_ERSTER_SCHUELER + 2)
    ]
    assert namen == ["Öztürk, Änne", "Straßer, Bernd"]


def test_xlsx_hat_leerspalten_zwischen_gruppen_und_vor_den_jahresnoten(session, graph):
    """Open point O-6."""
    zweite_gruppe = Notengruppe(
        kurs=graph.kurs,
        halbjahr=graph.halbjahr_1,
        bezeichnung="Mitarbeit",
        gewicht=Decimal("20"),
        reihenfolge=2,
    )
    session.add(zweite_gruppe)
    session.add(
        Leistung(
            notengruppe=zweite_gruppe,
            bezeichnung="Mitarbeit September",
            datum=date(2026, 9, 30),
            gewicht=Decimal("1.0"),
        )
    )
    session.commit()

    blattobjekt = geladen(als_xlsx(session))["BFS 26a Deutsch"]
    zeile = ZEILE_LEISTUNG
    # B: erste Leistung, C: leer, D: zweite Leistung, E: leer, F/G/H: Ergebnisse
    assert blattobjekt.cell(row=zeile, column=2).value.startswith("1. Klassenarbeit")
    assert blattobjekt.cell(row=zeile, column=3).value is None
    assert blattobjekt.cell(row=zeile, column=4).value.startswith("Mitarbeit September")
    assert blattobjekt.cell(row=zeile, column=5).value is None
    assert blattobjekt.cell(row=zeile, column=6).value == "1. Halbjahr"
    assert blattobjekt.cell(row=zeile, column=8).value == "Jahr"


def test_xlsx_werte_stimmen_mit_der_kursuebersicht_ueberein(session, graph):
    """Export and screen must never disagree."""
    bildschirm = blatt(session, graph.kurs, ANSICHT_JAHR)
    erwartet = bildschirm.zeilen[0].ergebnis_halbjahr[1]

    blattobjekt = geladen(als_xlsx(session))["BFS 26a Deutsch"]
    zelle = blattobjekt.cell(row=ZEILE_ERSTER_SCHUELER, column=4).value
    assert zelle == f"{erwartet.ganze_note} (2,0)"


def test_xlsx_ohne_kurse_ist_eine_gueltige_datei(session):
    mappe = geladen(als_xlsx(session))
    assert mappe.sheetnames == ["Kein Kurs"]


def test_xlsx_ohne_teilnehmer_bleibt_gueltig(session, graph):
    from app.services.verwaltung import setze_teilnahme

    for schueler in (graph.schueler_a, graph.schueler_b):
        setze_teilnahme(session, graph.kurs, schueler, False)
    session.commit()

    blattobjekt = geladen(als_xlsx(session))["BFS 26a Deutsch"]
    assert blattobjekt.cell(row=ZEILE_ERSTER_SCHUELER, column=1).value is None


def test_leistung_ohne_note_bleibt_leer_und_nicht_null(session, graph):
    blattobjekt = geladen(als_xlsx(session))["BFS 26a Deutsch"]
    # Bernd hat keine Note zur Klassenarbeit.
    assert blattobjekt.cell(row=ZEILE_ERSTER_SCHUELER + 1, column=2).value == OHNE_NOTE


# --------------------------------------------------------------------------
# Blattnamen: Excel ist hier eng
# --------------------------------------------------------------------------


def _kurs(bezeichnung: str, fach: str):
    class Vorgetaeuscht:
        pass

    kurs = Vorgetaeuscht()
    kurs.klasse = Vorgetaeuscht()
    kurs.klasse.bezeichnung = bezeichnung
    kurs.fach = fach
    return kurs


def test_blattname_wird_auf_einunddreissig_zeichen_gekuerzt():
    name = blattname(
        _kurs("Berufsfachschule 26a", "Berufsbezogener Unterricht"), set()
    )
    assert len(name) <= 31


def test_blattname_ersetzt_verbotene_zeichen():
    name = blattname(_kurs("BFS/26a", "Deutsch[1]"), set())
    for zeichen in "[]:*?/\\":
        assert zeichen not in name


def test_gleiche_blattnamen_werden_unterscheidbar():
    vergeben: set[str] = set()
    erster = blattname(_kurs("Berufsfachschule 26a", "Berufsbezogener U"), vergeben)
    zweiter = blattname(_kurs("Berufsfachschule 26a", "Berufsbezogener U"), vergeben)
    assert erster != zweiter
    assert len(zweiter) <= 31


def test_leerer_name_faellt_nicht_auf_die_nase():
    assert blattname(_kurs("", ""), set()) == "Kurs"


def test_nicht_gewertet_wird_im_export_gekennzeichnet(session, graph):
    session.add(
        Note(
            leistung=graph.leistung,
            schueler=graph.schueler_b,
            notenwert=None,
            status=NoteStatus.NICHT_GEWERTET,
        )
    )
    session.commit()

    blattobjekt = geladen(als_xlsx(session))["BFS 26a Deutsch"]
    assert blattobjekt.cell(row=ZEILE_ERSTER_SCHUELER + 1, column=2).value == NICHT_GEWERTET


def test_export_umfasst_alle_schuljahre_und_klassen(session, graph):
    """The operator's decision: everything, not one class at a time."""
    from app.db.models import Schuljahr

    altes = Schuljahr(
        bezeichnung="2025/26",
        beginn=date(2025, 8, 1),
        ende=date(2026, 7, 31),
        ist_aktiv=False,
    )
    alte_klasse = Klasse(schuljahr=altes, bezeichnung="BFS 25a")
    session.add(Kurs(klasse=alte_klasse, fach="Deutsch"))
    session.add(altes)
    session.commit()

    text = als_markdown(session)
    assert "## Schuljahr 2025/26" in text
    assert "## Schuljahr 2026/27" in text
    assert sorted(geladen(als_xlsx(session)).sheetnames) == [
        "BFS 25a Deutsch",
        "BFS 26a Deutsch",
    ]


def test_leere_notengruppe_bekommt_keine_spalte_wird_aber_genannt(session, graph):
    """A heading over nothing would also eat the spacer to the next group."""
    from app.services.export import ZEILE_GRUPPE, ZEILE_GRUPPENLISTE

    leer = Notengruppe(
        kurs=graph.kurs,
        halbjahr=graph.halbjahr_1,
        bezeichnung="Tests",
        gewicht=Decimal("30"),
        reihenfolge=2,
    )
    mitarbeit = Notengruppe(
        kurs=graph.kurs,
        halbjahr=graph.halbjahr_1,
        bezeichnung="Mitarbeit",
        gewicht=Decimal("20"),
        reihenfolge=3,
    )
    session.add_all([leer, mitarbeit])
    session.add(
        Leistung(
            notengruppe=mitarbeit,
            bezeichnung="Mitarbeit September",
            datum=date(2026, 9, 30),
            gewicht=Decimal("1.0"),
        )
    )
    session.commit()

    blattobjekt = geladen(als_xlsx(session))["BFS 26a Deutsch"]
    kopfzeilen = [
        blattobjekt.cell(row=ZEILE_GRUPPE, column=spalte).value for spalte in range(1, 7)
    ]
    # B Klassenarbeiten, C leer, D Mitarbeit, E leer -- Tests hat keine Spalte.
    assert kopfzeilen[1].startswith("Klassenarbeiten")
    assert kopfzeilen[2] is None
    assert kopfzeilen[3].startswith("Mitarbeit")
    assert kopfzeilen[4] is None

    # Verschwunden ist sie trotzdem nicht.
    liste = blattobjekt.cell(row=ZEILE_GRUPPENLISTE, column=1).value
    assert "Tests (Gewicht 30, ohne Leistungen)" in liste
