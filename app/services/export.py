"""Export of the whole dataset (specification 8, and section 11 against lock-in).

Both formats are built from :mod:`app.services.kursblatt` -- the same model
the course overview on screen is built from. That is the point, not a
convenience: two separate ways of computing the same grades drift apart
sooner or later, and it only shows when somebody compares the numbers.

**Deviation from 8, consistent with the screen: no group averages.** They are
not a step of the single-stage calculation, so a column for them would invite
arithmetic that does not add up.

Everything is produced in memory. Nothing is written to disk, nothing is
exported on a schedule -- the file exists only as the answer to a request
somebody made (specification 8, 11).
"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.db.models import Kurs
from app.enums import NoteStatus
from app.grading.notenwert import als_anzeige, als_dezimalanzeige
from app.services.kursblatt import ANSICHT_JAHR, Kursblatt, blatt
from app.services.settings import lies_einstellungen
from app.services.sorting import anzeigename, vereinfacht

OHNE_NOTE = "—"
NICHT_ERBRACHT = "n.e."
NICHT_GEWERTET = "n.g."

LEGENDE = (
    f"{NICHT_ERBRACHT} = nicht erbracht (zählt als 6) · "
    f"{NICHT_GEWERTET} = nicht gewertet (zählt nicht) · "
    f"{OHNE_NOTE} = keine Note eingetragen"
)

# Excel refuses these in a sheet title, and titles over 31 characters are only
# warned about -- some applications then cannot read the file at all.
VERBOTENE_ZEICHEN = str.maketrans({zeichen: "-" for zeichen in "[]:*?/\\"})
MAX_BLATTNAME = 31


def alle_kurse(session: Session) -> list[Kurs]:
    """Every course of every class of every school year, in a stable order."""
    return sorted(
        session.query(Kurs).all(),
        key=lambda kurs: (
            kurs.klasse.schuljahr.beginn,
            vereinfacht(kurs.klasse.bezeichnung),
            vereinfacht(kurs.fach),
            kurs.id,
        ),
    )


def notenzelle(note) -> str:
    """One grade as it is read: 0.7 is "1+", a status is named, nothing is 0."""
    if note is None:
        return OHNE_NOTE
    if note.status == NoteStatus.NICHT_ERBRACHT:
        return NICHT_ERBRACHT
    if note.status == NoteStatus.NICHT_GEWERTET:
        return NICHT_GEWERTET
    return als_anzeige(note.notenwert)


def ergebniszelle(ergebnis, festsetzung) -> str:
    """Whole grade with the calculated value beside it, as on screen."""
    if festsetzung is not None and ergebnis is not None:
        return f"{als_anzeige(festsetzung.notenwert)} ({als_dezimalanzeige(ergebnis.wert)})"
    if festsetzung is not None:
        return f"{als_anzeige(festsetzung.notenwert)} (festgesetzt)"
    if ergebnis is not None:
        return f"{ergebnis.ganze_note} ({als_dezimalanzeige(ergebnis.wert)})"
    return OHNE_NOTE


def blattname(kurs: Kurs, vergeben: set[str]) -> str:
    """Sheet title within Excel's limits, and unique within the workbook."""
    roh = f"{kurs.klasse.bezeichnung} {kurs.fach}".translate(VERBOTENE_ZEICHEN)
    sauber = roh.strip().strip("'") or "Kurs"

    name = sauber[:MAX_BLATTNAME]
    zaehler = 2
    while name in vergeben:
        anhang = f" ({zaehler})"
        name = sauber[: MAX_BLATTNAME - len(anhang)].rstrip() + anhang
        zaehler += 1
    vergeben.add(name)
    return name


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def _kursabschnitt(kursblatt: Kursblatt, namensanzeige) -> list[str]:
    kurs = kursblatt.kurs
    zeilen = [f"### {kurs.klasse.bezeichnung} — {kurs.fach}", ""]

    if not kursblatt.zeilen:
        zeilen += ["Keine aktiven Teilnehmer.", ""]
        return zeilen

    gruppen = [
        f"{g.halbjahr.nummer}. HJ · {g.notengruppe.bezeichnung} "
        f"(Gewicht {g.notengruppe.gewicht})"
        for g in kursblatt.gruppen
    ]
    if gruppen:
        zeilen += ["Notengruppen: " + " · ".join(gruppen), ""]

    kopf = ["Schüler"]
    for gruppe in kursblatt.gruppen:
        for leistung in gruppe.leistungen:
            kopf.append(
                f"{gruppe.halbjahr.nummer}. HJ · {gruppe.notengruppe.bezeichnung} · "
                f"{leistung.bezeichnung} ({leistung.datum:%d.%m.%Y})"
            )
    kopf += ["1. Halbjahr", "2. Halbjahr", "Jahr"]

    zeilen.append("| " + " | ".join(kopf) + " |")
    zeilen.append("|" + "---|" * len(kopf))

    for zeile in kursblatt.zeilen:
        felder = [anzeigename(zeile.schueler.vorname, zeile.schueler.nachname, namensanzeige)]
        for gruppe in kursblatt.gruppen:
            for leistung in gruppe.leistungen:
                felder.append(notenzelle(zeile.noten.get(leistung.id)))
        for nummer in (1, 2):
            felder.append(
                ergebniszelle(
                    zeile.ergebnis_halbjahr.get(nummer),
                    zeile.festsetzung_halbjahr.get(nummer),
                )
            )
        felder.append(ergebniszelle(zeile.jahresergebnis, zeile.jahresfestsetzung))
        zeilen.append("| " + " | ".join(felder) + " |")

    zeilen.append("")
    return zeilen


def als_markdown(session: Session) -> str:
    """The whole dataset as structured text, one section per course (8)."""
    namensanzeige = lies_einstellungen(session).namensanzeige
    zeilen = [
        f"# Notenverwaltung — Export vom {date.today():%d.%m.%Y}",
        "",
        "Dieser Export enthält Klarnamen.",
        "",
        LEGENDE,
        "",
    ]

    letztes_schuljahr = None
    for kurs in alle_kurse(session):
        schuljahr = kurs.klasse.schuljahr
        if schuljahr is not letztes_schuljahr:
            zeilen += [f"## Schuljahr {schuljahr.bezeichnung}", ""]
            letztes_schuljahr = schuljahr
        zeilen += _kursabschnitt(blatt(session, kurs, ANSICHT_JAHR), namensanzeige)

    if letztes_schuljahr is None:
        zeilen += ["Es ist noch kein Kurs angelegt.", ""]

    return "\n".join(zeilen)


# ---------------------------------------------------------------------------
# XLSX, layout per open point O-6
# ---------------------------------------------------------------------------

ZEILE_KOPF = 1
ZEILE_UNTERTITEL = 2
ZEILE_GRUPPENLISTE = 3
ZEILE_HALBJAHR = 4
ZEILE_GRUPPE = 5
ZEILE_LEISTUNG = 6
ZEILE_ERSTER_SCHUELER = 7

GEDREHT = Alignment(textRotation=90, vertical="bottom", horizontal="center")
FETT = Font(bold=True)


def _beschriftung(blattobjekt, zeile: int, spalte: int, text: str) -> None:
    zelle = blattobjekt.cell(row=zeile, column=spalte, value=text)
    zelle.alignment = GEDREHT


def _kursblatt_schreiben(mappe: Workbook, kursblatt: Kursblatt, name: str, namensanzeige):
    kurs = kursblatt.kurs
    blattobjekt = mappe.create_sheet(title=name)

    blattobjekt.cell(
        row=ZEILE_KOPF, column=1, value=f"{kurs.klasse.bezeichnung} — {kurs.fach}"
    ).font = FETT
    blattobjekt.cell(
        row=ZEILE_UNTERTITEL,
        column=1,
        value=(
            f"Schuljahr {kurs.klasse.schuljahr.bezeichnung} · "
            f"erstellt am {date.today():%d.%m.%Y} · {LEGENDE}"
        ),
    )

    # A group without assessments gets no column: it would be a heading over
    # nothing and would eat the spacer to the next group. Its weight is still
    # named in the list above, so nothing is hidden.
    gruppen = [gruppe for gruppe in kursblatt.gruppen if gruppe.leistungen]

    alle_gruppen = " · ".join(
        f"{g.halbjahr.nummer}. HJ {g.notengruppe.bezeichnung} "
        f"(Gewicht {g.notengruppe.gewicht}"
        + (", ohne Leistungen)" if not g.leistungen else ")")
        for g in kursblatt.gruppen
    )
    if alle_gruppen:
        blattobjekt.cell(
            row=ZEILE_GRUPPENLISTE, column=1, value="Notengruppen: " + alle_gruppen
        )

    # Column 1 holds the names; the assessment columns start at 2, with an
    # empty column between two groups and before the year grades (O-6).
    spalte = 2
    spalten_je_leistung: dict[int, int] = {}
    for index, gruppe in enumerate(gruppen):
        if index > 0:
            spalte += 1  # Abstandsspalte
        _beschriftung(
            blattobjekt, ZEILE_HALBJAHR, spalte, f"{gruppe.halbjahr.nummer}. Halbjahr"
        )
        _beschriftung(
            blattobjekt,
            ZEILE_GRUPPE,
            spalte,
            f"{gruppe.notengruppe.bezeichnung} (Gewicht {gruppe.notengruppe.gewicht})",
        )
        for leistung in gruppe.leistungen:
            _beschriftung(
                blattobjekt,
                ZEILE_LEISTUNG,
                spalte,
                f"{leistung.bezeichnung} ({leistung.datum:%d.%m.%Y})",
            )
            spalten_je_leistung[leistung.id] = spalte
            spalte += 1

    if gruppen:
        spalte += 1  # Abstand vor den Jahresnoten

    erste_ergebnisspalte = spalte
    for versatz, beschriftung in enumerate(("1. Halbjahr", "2. Halbjahr", "Jahr")):
        _beschriftung(blattobjekt, ZEILE_LEISTUNG, erste_ergebnisspalte + versatz, beschriftung)

    blattobjekt.cell(row=ZEILE_LEISTUNG, column=1, value="Schüler").font = FETT
    for versatz, zeile in enumerate(kursblatt.zeilen):
        nummer = ZEILE_ERSTER_SCHUELER + versatz
        blattobjekt.cell(
            row=nummer,
            column=1,
            value=anzeigename(
                zeile.schueler.vorname, zeile.schueler.nachname, namensanzeige
            ),
        )
        for leistung_id, spaltennummer in spalten_je_leistung.items():
            blattobjekt.cell(
                row=nummer,
                column=spaltennummer,
                value=notenzelle(zeile.noten.get(leistung_id)),
            )
        for hj in (1, 2):
            blattobjekt.cell(
                row=nummer,
                column=erste_ergebnisspalte + hj - 1,
                value=ergebniszelle(
                    zeile.ergebnis_halbjahr.get(hj), zeile.festsetzung_halbjahr.get(hj)
                ),
            )
        blattobjekt.cell(
            row=nummer,
            column=erste_ergebnisspalte + 2,
            value=ergebniszelle(zeile.jahresergebnis, zeile.jahresfestsetzung),
        )

    blattobjekt.column_dimensions["A"].width = 28
    for nummer in range(2, erste_ergebnisspalte + 3):
        blattobjekt.column_dimensions[get_column_letter(nummer)].width = 6
    blattobjekt.row_dimensions[ZEILE_LEISTUNG].height = 130
    blattobjekt.row_dimensions[ZEILE_GRUPPE].height = 110
    blattobjekt.row_dimensions[ZEILE_HALBJAHR].height = 80
    # Names and labels stay put while the sheet scrolls.
    blattobjekt.freeze_panes = f"B{ZEILE_ERSTER_SCHUELER}"


def als_xlsx(session: Session) -> bytes:
    """The whole dataset as a workbook, one sheet per course (8)."""
    namensanzeige = lies_einstellungen(session).namensanzeige
    mappe = Workbook()
    mappe.remove(mappe.active)

    vergeben: set[str] = set()
    for kurs in alle_kurse(session):
        _kursblatt_schreiben(
            mappe, blatt(session, kurs, ANSICHT_JAHR), blattname(kurs, vergeben),
            namensanzeige,
        )

    if not mappe.sheetnames:
        leer = mappe.create_sheet(title="Kein Kurs")
        leer.cell(row=1, column=1, value="Es ist noch kein Kurs angelegt.")

    puffer = BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()
